#!/usr/bin/env python3
"""Reconnaissance for the Queensland workforce data, run on a GitHub runner.

WHY A RUNNER. The current State of the Sector workbooks live on
www.data.qld.gov.au, which answers `x-amzn-waf-action: challenge` and returns a
JavaScript interstitial instead of the file. The sandbox this repo is developed
in cannot execute that challenge, so the file has never been seen. Everything
the CKAN datastore will serve stops at March 2023, and the older profiles on
forgov.qld.gov.au download fine — so it is that one host, not the jurisdiction.

WHAT THIS DOES, AND WHY IT IS NOT THE PARSER. It answers two questions that
have to be answered before a parser can honestly be written:

  1. Does a GitHub runner get the file, or the same challenge?
  2. What is actually in it? The current product is the "State of the sector
     report" workbook, which is not the "biannual workforce profile" the
     datastore holds — a parser written against the old shape would be a guess.

So it prints the structure to the job log and stops. The parser goes into
scripts/gen-gov-workforce.py once there is a measurement to write it against.

    python3 scripts/qld-workforce-probe.py
"""
import io, json, re, sys, urllib.error, urllib.request

API = 'https://data.qld.gov.au/api/3/action'
DATASET = 'queensland-public-service-workforce-quarterly-profile'
UA = ('Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36')
HEADERS = {
    'User-Agent': UA,
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
    'Accept-Language': 'en-AU,en;q=0.9',
    'Referer': f'https://www.data.qld.gov.au/dataset/{DATASET}',
}


def fetch(url, binary=False):
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read() if binary else r.read().decode('utf-8', 'replace')


def main():
    pkg = json.loads(fetch(f'{API}/package_show?id={DATASET}'))['result']
    # Newest "State of the sector" workbook — that is the current product.
    cands = [r for r in pkg['resources']
             if 'state of the sector' in r['name'].lower()
             and (r.get('format') or '').lower() in ('xlsx', 'xls')]
    cands.sort(key=lambda r: re.search(r'(20\d\d)', r['name']).group(1)
               if re.search(r'(20\d\d)', r['name']) else '', reverse=True)
    if not cands:
        print('NO state-of-the-sector workbook in the dataset', file=sys.stderr)
        return 1
    res = cands[0]
    print(f'resource : {res["name"]}')
    print(f'url      : {res["url"]}')

    try:
        raw = fetch(res['url'], binary=True)
    except urllib.error.HTTPError as e:
        print(f'FETCH FAILED: HTTP {e.code} — the runner is challenged too')
        return 1

    print(f'bytes    : {len(raw)}')
    if raw[:2] != b'PK':
        # A WAF interstitial is HTML, and small.
        print('NOT A WORKBOOK — first 300 bytes follow. The runner was challenged.')
        print(raw[:300].decode('utf-8', 'replace'))
        return 1

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    print(f'sheets   : {len(wb.sheetnames)}')
    for n in wb.sheetnames:
        print(f'  - {n}')

    # For each sheet, print the first rows that carry text, so the agency table
    # can be identified by eye rather than by a rule written in advance.
    for n in wb.sheetnames:
        ws = wb[n]
        print(f'\n===== {n} =====')
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True)):
            cells = ['' if c is None else str(c).strip() for c in (row or ())[:10]]
            if any(cells):
                print(f'  {i+1:3d} | ' + ' | '.join(c[:28] for c in cells))
    return 0


if __name__ == '__main__':
    sys.exit(main())
